"""Turn a resume plus a short intake into a "circle of competence" profile.

The idea (Peter Lynch's "invest in what you know", Munger's circle of competence):
the industries you work in, read about, and spend money in are where you are most
likely to notice change before the market does. The profile records *why* each
sector is in your circle so every idea can cite the evidence.

A caution baked into the output: research by Long, Fernbach & De Langhe (2018,
Journal of Marketing Research) found that feeling you understand a company makes
people rate it as less risky, with no link to its actual risk. So familiarity
raises a name's *ranking*, never its position size or risk budget.
"""

from __future__ import annotations

import io
import re
from collections import defaultdict
from dataclasses import asdict, dataclass, field

from .universe import DOMAIN_BY_KEY, DOMAINS, TICKERS, lookup

# Source weights. What you pay attention to by choice (reading, spending) is
# weighted a bit above what's merely on a resume; employers count the most
# because you saw the business from the inside.
W_RESUME = 1.0
W_INTERESTS = 1.3
W_EMPLOYER = 4.0
W_BRAND = 2.0
W_EXPLICIT_TICKER = 3.0

RISK_LABELS = {
    1: "Learn (shares only)",
    2: "Conservative",
    3: "Moderate",
    4: "Growth",
    5: "Speculative",
}


@dataclass
class DomainHit:
    key: str
    label: str
    score: float = 0.0
    familiarity: int = 0  # 0-100
    evidence: list[str] = field(default_factory=list)


@dataclass
class Profile:
    name: str = ""
    years_experience: int | None = None
    capital: float = 10000.0
    risk_level: int = 2
    holdings: dict[str, int] = field(default_factory=dict)  # symbol -> shares
    domains: list[dict] = field(default_factory=list)
    tickers: list[dict] = field(default_factory=list)  # {symbol, name, domain, familiarity, why}
    watchlist: list[str] = field(default_factory=list)
    summary: str = ""
    caution: str = ""
    source: str = "rules"

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> "Profile":
        known = {k: d[k] for k in cls.__dataclass_fields__ if k in d}
        return cls(**known)

    def familiarity_for(self, symbol: str) -> tuple[int, list[str]]:
        for t in self.tickers:
            if t["symbol"] == symbol:
                return t["familiarity"], t["why"]
        info = lookup(symbol)
        if info:
            for d in self.domains:
                if d["key"] == info.domain:
                    return int(d["familiarity"] * 0.7), [f"In your {d['label']} circle"]
        return 0, []


# --------------------------------------------------------------------------- text

# Upload limits (CHL-PS3): a resume is a few pages, so anything far bigger is refused before parsing.
MAX_ZIP_MEMBERS = 200
MAX_ZIP_UNPACKED = 20_000_000
MAX_ZIP_RATIO = 100
MAX_PDF_PAGES = 30
MAX_TEXT_CHARS = 100_000
# CHL-PS14: pypdf parses every content-stream operator in Python, so a tiny file whose stream inflates to
# tens of MB burns minutes and GBs. A real resume's pages decode to tens of KB, so 2 MB in all is plenty.
MAX_PDF_CONTENT = 2_000_000
_PDF_DECODE_LIMIT = 10_000_000  # pypdf's own per-stream inflate caps, lowered from its 75 MB default


class FileTooComplex(ValueError):
    """The upload unpacks to far more than a resume would (a decompression bomb, or simply huge)."""


def _check_zip(data: bytes) -> None:
    import zipfile
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        infos = z.infolist()
        unpacked = sum(i.file_size for i in infos)
        if len(infos) > MAX_ZIP_MEMBERS:
            raise FileTooComplex(f"{len(infos)} parts in the file (max {MAX_ZIP_MEMBERS})")
        if unpacked > MAX_ZIP_UNPACKED:
            raise FileTooComplex(f"unpacks to {unpacked:,} bytes (max {MAX_ZIP_UNPACKED:,})")
        if unpacked > MAX_ZIP_RATIO * max(len(data), 1):
            raise FileTooComplex(f"unpacks to over {MAX_ZIP_RATIO}x its size")


def _pdf_stream_size(stream, budget: int) -> int:
    """Decoded size of one PDF stream, inflating at most budget + 1 bytes for plain Flate streams."""
    import zlib
    filters = stream.get("/Filter")
    filters = [] if filters is None else [str(f) for f in (filters if isinstance(filters, list) else [filters])]
    raw = getattr(stream, "_data", None)
    if raw is not None and not filters:
        return len(raw)
    if raw is not None and filters in (["/FlateDecode"], ["/Fl"]):
        try:
            return len(zlib.decompressobj().decompress(raw, budget + 1))
        except zlib.error:
            pass  # damaged stream: let pypdf (and its recovery) measure it below
    try:  # other filter chains: pypdf decodes under its lowered limits
        return len(stream.get_data())
    except Exception as e:
        raise FileTooComplex(f"a PDF stream couldn't be decoded within limits ({type(e).__name__})")


def _check_pdf(pages) -> None:
    """Refuse a PDF whose page content (and the forms the pages draw) decodes to over MAX_PDF_CONTENT
    bytes, before any text extraction. A stream shared by many pages counts once per page, since pypdf
    parses it once per page."""
    import pypdf.filters as pf
    for name in ("ZLIB_MAX_OUTPUT_LENGTH", "LZW_MAX_OUTPUT_LENGTH", "RUN_LENGTH_MAX_OUTPUT_LENGTH"):
        if getattr(pf, name, 0) > _PDF_DECODE_LIMIT:
            setattr(pf, name, _PDF_DECODE_LIMIT)
    total = 0

    def add(stream):
        nonlocal total
        total += _pdf_stream_size(stream, MAX_PDF_CONTENT - total)
        if total > MAX_PDF_CONTENT:
            raise FileTooComplex(f"page content decodes to over {MAX_PDF_CONTENT:,} bytes")

    def forms(resources, seen, depth):
        xobjects = (resources or {}).get("/XObject") if hasattr(resources, "get") else None
        if depth > 5 or not xobjects:
            return
        for ref in xobjects.get_object().values():
            obj = ref.get_object()
            if obj.get("/Subtype") != "/Form" or id(obj) in seen:
                continue
            seen.add(id(obj))
            add(obj)
            forms(obj.get("/Resources"), seen, depth + 1)

    for page in pages:
        contents = page.get("/Contents")
        if contents is not None:
            contents = contents.get_object()
            for part in (contents if isinstance(contents, list) else [contents]):
                add(part.get_object())
        res = page.get("/Resources")
        forms(res.get_object() if res is not None else None, set(), 0)


def _capped(parts) -> str:
    """Join text parts, stopping once MAX_TEXT_CHARS is reached."""
    out, size = [], 0
    for part in parts:
        out.append(part)
        size += len(part) + 1
        if size >= MAX_TEXT_CHARS:
            break
    return "\n".join(out)[:MAX_TEXT_CHARS]


def extract_text(filename: str, data: bytes) -> str:
    """Pull plain text out of a PDF, DOCX, or text upload: at most 30 PDF pages and 100k characters.
    Raises FileTooComplex for a .docx, or a PDF's page content, that would unpack to far more than a resume."""
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:4] == b"%PDF":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = [reader.pages[i] for i in range(min(len(reader.pages), MAX_PDF_PAGES))]
        _check_pdf(pages)  # before pypdf parses a single operator (CHL-PS14)
        return _capped((p.extract_text() or "") for p in pages)
    if name.endswith(".docx") or data[:2] == b"PK":
        _check_zip(data)  # before python-docx unpacks anything
        import docx
        doc = docx.Document(io.BytesIO(data))

        def parts():
            yield from (p.text for p in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    yield from (cell.text for cell in row.cells)
        return _capped(parts())
    return data[:MAX_TEXT_CHARS * 4].decode("utf-8", errors="ignore")[:MAX_TEXT_CHARS]


def _norm(text: str) -> str:
    return " " + re.sub(r"\s+", " ", text.lower()) + " "


def _count(term: str, text: str) -> int:
    pattern = r"(?<![a-z0-9])" + re.escape(term.lower().strip()) + r"(?![a-z0-9])"
    return len(re.findall(pattern, text))


_YEAR_RANGE = re.compile(r"((?:19|20)\d{2})\s*(?:-|–|—|to)\s*((?:19|20)\d{2}|present|current|now)", re.I)
_YEARS_EXP = re.compile(r"(\d{1,2})\+?\s*years?\s+(?:of\s+)?experience", re.I)


def estimate_years(text: str, this_year: int) -> int | None:
    m = _YEARS_EXP.search(text)
    if m:
        return int(m.group(1))
    spans = []
    for a, b in _YEAR_RANGE.findall(text):
        start = int(a)
        end = this_year if not b[:1].isdigit() else int(b)
        if 1970 <= start <= end <= this_year:
            spans.append((start, end))
    if not spans:
        return None
    return max(0, max(e for _, e in spans) - min(s for s, _ in spans))


_TICKER_TOKEN = re.compile(r"\$?\b([A-Z]{1,5}(?:\.[A-Z])?)\b")


def explicit_tickers(text: str) -> list[str]:
    found = []
    for tok in _TICKER_TOKEN.findall(text or ""):
        if len(tok) >= 2 and lookup(tok) and tok not in found:
            found.append(lookup(tok).symbol)
    return found


# ------------------------------------------------------------------------ builder

def build_profile(
    resume_text: str = "",
    interests: str = "",
    brands: str = "",
    holdings: dict[str, int] | None = None,
    capital: float = 10000.0,
    risk_level: int = 2,
    name: str = "",
    this_year: int = 2026,
    domain_boosts: dict[str, tuple[float, str]] | None = None,
    ticker_boosts: dict[str, str] | None = None,
) -> Profile:
    """Build a profile. `domain_boosts`/`ticker_boosts` carry optional extra evidence
    (e.g. from an LLM read of the resume): {domain: (points, evidence)}, {symbol: reason}."""
    resume_n = _norm(resume_text)
    interests_n = _norm(interests)
    brands_n = _norm(brands)

    hits: dict[str, DomainHit] = {d.key: DomainHit(d.key, d.label) for d in DOMAINS}
    ticker_score: dict[str, float] = defaultdict(float)
    ticker_why: dict[str, list[str]] = defaultdict(list)

    # 1. Keywords -> domains
    for d in DOMAINS:
        for kw in d.keywords:
            r = _count(kw, resume_n)
            i = _count(kw, interests_n)
            if r:
                hits[d.key].score += W_RESUME * min(r, 4)
                hits[d.key].evidence.append(f"resume mentions \"{kw.strip()}\"")
            if i:
                hits[d.key].score += W_INTERESTS * min(i, 3)
                hits[d.key].evidence.append(f"you follow \"{kw.strip()}\"")

    # 2. Company names: employers in the resume, brands you use, things you read
    for t in TICKERS:
        names = (t.name.lower(),) + t.aliases
        r_alias = next((a for a in names if _count(a, resume_n)), None)
        r = 1 if r_alias else 0
        b = max((_count(a, brands_n) for a in names), default=0)
        i = max((_count(a, interests_n) for a in names), default=0)
        if r:
            ticker_score[t.symbol] += W_EMPLOYER
            shown = t.name if r_alias == t.name.lower() else f"{r_alias.title()} ({t.name})"
            ticker_why[t.symbol].append(f"your resume mentions {shown}")
            hits[t.domain].score += W_EMPLOYER / 2
            hits[t.domain].evidence.append(f"resume mentions {shown}")
        if b:
            ticker_score[t.symbol] += W_BRAND
            ticker_why[t.symbol].append(f"you're a customer of {t.name}")
            hits[t.domain].score += W_BRAND / 2
        if i:
            ticker_score[t.symbol] += W_INTERESTS * 1.5
            ticker_why[t.symbol].append(f"you read/follow {t.name}")
            hits[t.domain].score += W_INTERESTS

    for sym in explicit_tickers(interests + " " + brands):
        ticker_score[sym] += W_EXPLICIT_TICKER
        ticker_why[sym].append("you named this ticker")
        hits[lookup(sym).domain].score += 1.0

    for key, (pts, ev) in (domain_boosts or {}).items():
        if key in hits:
            hits[key].score += float(pts)
            hits[key].evidence.append(ev)
    for sym, reason in (ticker_boosts or {}).items():
        info = lookup(sym)
        if info:
            ticker_score[info.symbol] += W_BRAND
            ticker_why[info.symbol].append(reason)

    holdings = {k.upper(): int(v) for k, v in (holdings or {}).items() if int(v) > 0}
    for sym in holdings:
        info = lookup(sym)
        if info:
            ticker_score[info.symbol] += W_EXPLICIT_TICKER
            ticker_why[info.symbol].append(f"you own {holdings[sym]} shares")

    # 3. Normalize domain familiarity (top domain = 100, log-ish damping)
    ranked = sorted((h for h in hits.values() if h.score > 0), key=lambda h: -h.score)
    top = ranked[0].score if ranked else 1.0
    for h in ranked:
        h.familiarity = max(5, min(100, round(100 * (h.score / top) ** 0.7)))
        # de-dup evidence, keep it short
        seen, ev = set(), []
        for e in h.evidence:
            if e not in seen:
                seen.add(e)
                ev.append(e)
        h.evidence = ev[:6]

    domains = [
        {"key": h.key, "label": h.label, "familiarity": h.familiarity, "evidence": h.evidence,
         "etf": DOMAIN_BY_KEY[h.key].etf, "blurb": DOMAIN_BY_KEY[h.key].blurb}
        for h in ranked[:6]
    ]
    if not domains:
        # Nothing matched: start from broad, durable franchises and say so.
        for key in ("payments_fintech", "restaurants_food", "retail_consumer"):
            d = DOMAIN_BY_KEY[key]
            domains.append({"key": key, "label": d.label, "familiarity": 20,
                            "evidence": ["default starter circle (add more detail to personalize)"],
                            "etf": d.etf, "blurb": d.blurb})

    dom_fam = {d["key"]: d["familiarity"] for d in domains}

    # 4. Tickers: everything in your top domains, boosted by direct evidence
    tickers: dict[str, dict] = {}
    for d in domains[:4]:
        for sym in DOMAIN_BY_KEY[d["key"]].tickers:
            info = lookup(sym)
            tickers[sym] = {"symbol": sym, "name": info.name, "domain": info.domain,
                            "familiarity": int(d["familiarity"] * 0.75),
                            "why": [f"core name in your {d['label']} circle"]}
    for sym, sc in ticker_score.items():
        info = lookup(sym)
        base = int(dom_fam.get(info.domain, 0) * 0.75)
        fam = min(100, base + int(12 * sc))
        why = ticker_why[sym] + ([f"in your {DOMAIN_BY_KEY[info.domain].label} circle"] if base else [])
        tickers[sym] = {"symbol": sym, "name": info.name, "domain": info.domain,
                        "familiarity": fam, "why": why[:4]}
    ticker_list = sorted(tickers.values(), key=lambda t: -t["familiarity"])

    years = estimate_years(resume_text, this_year)
    top_labels = ", ".join(d["label"] for d in domains[:3])
    summary = (f"Your circle of competence centers on {top_labels}."
               + (f" About {years} years of work history detected." if years else ""))
    caution = ("Knowing an industry is an edge for finding ideas, not proof they're safe. "
               "Research shows familiarity makes companies feel less risky than they are, so "
               "Charlie sizes every idea by its objective risk, not by how well you know it.")

    return Profile(
        name=name, years_experience=years, capital=float(capital),
        risk_level=max(1, min(5, int(risk_level))), holdings=holdings,
        domains=domains, tickers=ticker_list[:20],
        watchlist=[t["symbol"] for t in ticker_list[:12]],
        summary=summary, caution=caution,
    )
